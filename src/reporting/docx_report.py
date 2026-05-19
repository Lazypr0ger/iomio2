from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from src.config import REPORTS_DIR
from src.functions.base import ObjectiveFunction
from src.models import OptimizationResult
from src.reporting.tables import (
    result_history_to_dataframe,
    results_to_summary_dataframe,
)
from src.visualization.contours import plot_contour_with_tracks
from src.visualization.convergence import (
    plot_function_decrease,
    plot_gradient_norm_convergence,
)
from src.visualization.polynomial_sections import plot_polynomial_2d_section


FUNCTION_RU_NAMES = {
    "Himmelblau": "Химмельблау",
    "Rosenbrock": "Розенброк",
    "Booth": "Бут",
    "Beale": "Бил",
    "Matyas": "Матьяс",
    "ThreeHumpCamel": "Трёхгорбый верблюд",
    "Sphere": "Сфера",
    "Rastrigin": "Растригин",
    "Polynomial10DVariant3": "Полином 10D, вариант 3",
    "CustomPolynomial10D": "Произвольный полином 10D",
}


METHOD_RU_NAMES = {
    "CG Fletcher-Reeves": "Сопряжённые градиенты Флетчера–Ривза",
    "Modified Newton": "Модифицированный метод Ньютона",
    "BFGS": "BFGS",
}


SUMMARY_COLUMNS_RU = {
    "function": "Функция",
    "method": "Метод",
    "converged": "Сошёлся",
    "iterations": "Итерации",
    "function_evaluations": "Вычисления f",
    "gradient_evaluations": "Вычисления ∇f",
    "hessian_evaluations": "Вычисления H",
    "x_star": "x*",
    "f_star": "f(x*)",
    "grad_norm": "||∇f(x*)||",
    "message": "Сообщение",
}


HISTORY_COLUMNS_RU = {
    "k": "k",
    "f_value": "f(xᵏ)",
    "grad_norm": "||∇f(xᵏ)||",
    "step_size": "λᵏ",
    "x1": "x₁ᵏ",
    "x2": "x₂ᵏ",
    "x": "xᵏ",
    "step_norm": "||xᵏ - xᵏ⁻¹||",
    "beta": "βᵏ",
    "restart": "Рестарт",
    "regularization": "Регуляризация",
    "rho": "ρ",
    "update": "Обновление",
}


def get_function_ru_name(name: str) -> str:
    return FUNCTION_RU_NAMES.get(name, name)


def get_method_ru_name(name: str) -> str:
    return METHOD_RU_NAMES.get(name, name)


def format_array(array: np.ndarray, precision: int = 6) -> str:
    return np.array2string(
        array,
        precision=precision,
        separator=", ",
        suppress_small=True,
    )


def format_value(value: object) -> str:
    if value is None:
        return "—"

    if isinstance(value, bool):
        return "Да" if value else "Нет"

    if isinstance(value, float):
        return f"{value:.6g}"

    if isinstance(value, np.floating):
        return f"{float(value):.6g}"

    if isinstance(value, np.integer):
        return str(int(value))

    return str(value)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_paragraph(document: Document, text: str = "") -> None:
    paragraph = document.add_paragraph(text)
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(6)


def add_centered_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(16)


def add_dataframe_table(
    document: Document,
    dataframe: pd.DataFrame,
    max_rows: int | None = None,
) -> None:
    if dataframe.empty:
        add_paragraph(document, "Нет данных для таблицы.")
        return

    table_dataframe = dataframe.copy()

    if max_rows is not None and len(table_dataframe) > max_rows:
        head = table_dataframe.head(max_rows // 2)
        tail = table_dataframe.tail(max_rows - len(head))
        separator = pd.DataFrame(
            [
                {
                    column: "..."
                    for column in table_dataframe.columns
                }
            ]
        )
        table_dataframe = pd.concat([head, separator, tail], ignore_index=True)

    table = document.add_table(
        rows=1,
        cols=len(table_dataframe.columns),
    )
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for index, column_name in enumerate(table_dataframe.columns):
        header_cells[index].text = str(column_name)

    for _, row in table_dataframe.iterrows():
        cells = table.add_row().cells

        for index, value in enumerate(row):
            cells[index].text = format_value(value)


def add_image_if_exists(
    document: Document,
    image_path: Path,
    width_inches: float = 6.5,
) -> None:
    if not image_path.exists():
        add_paragraph(document, f"График не найден: {image_path}")
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def prepare_summary_dataframe(results: list[OptimizationResult]) -> pd.DataFrame:
    dataframe = results_to_summary_dataframe(results)

    if "function" in dataframe.columns:
        dataframe["function"] = dataframe["function"].map(get_function_ru_name)

    if "method" in dataframe.columns:
        dataframe["method"] = dataframe["method"].map(get_method_ru_name)

    columns_to_keep = [
        "function",
        "method",
        "converged",
        "iterations",
        "function_evaluations",
        "gradient_evaluations",
        "hessian_evaluations",
        "x_star",
        "f_star",
        "grad_norm",
    ]

    dataframe = dataframe[columns_to_keep]
    dataframe = dataframe.rename(columns=SUMMARY_COLUMNS_RU)

    return dataframe


def prepare_function_summary_dataframe(
    results: list[OptimizationResult],
) -> pd.DataFrame:
    dataframe = prepare_summary_dataframe(results)

    columns_to_keep = [
        "Метод",
        "Итерации",
        "Вычисления f",
        "Вычисления ∇f",
        "Вычисления H",
        "x*",
        "f(x*)",
        "||∇f(x*)||",
    ]

    return dataframe[columns_to_keep]


def prepare_history_dataframe(
    result: OptimizationResult,
    max_rows: int = 12,
) -> pd.DataFrame:
    dataframe = result_history_to_dataframe(result)

    if result.x_star.size == 2:
        columns_to_keep = [
            column
            for column in [
                "k",
                "x1",
                "x2",
                "f_value",
                "grad_norm",
                "step_size",
                "beta",
                "regularization",
                "rho",
                "update",
            ]
            if column in dataframe.columns
        ]
    else:
        columns_to_keep = [
            column
            for column in [
                "k",
                "f_value",
                "grad_norm",
                "step_size",
                "step_norm",
                "beta",
                "regularization",
                "rho",
                "update",
            ]
            if column in dataframe.columns
        ]

    dataframe = dataframe[columns_to_keep]
    dataframe = dataframe.rename(columns=HISTORY_COLUMNS_RU)

    if len(dataframe) > max_rows:
        head = dataframe.head(max_rows // 2)
        tail = dataframe.tail(max_rows - len(head))
        separator = pd.DataFrame(
            [
                {
                    column: "..."
                    for column in dataframe.columns
                }
            ]
        )
        dataframe = pd.concat([head, separator, tail], ignore_index=True)

    return dataframe


def get_results_for_function(
    function: ObjectiveFunction,
    results: list[OptimizationResult],
) -> list[OptimizationResult]:
    return [
        result
        for result in results
        if result.function_name == function.name
    ]


def add_title_page(document: Document) -> None:
    add_centered_title(document, "Лабораторная работа №2")
    add_centered_title(document, "Многомерная нелинейная оптимизация")

    add_paragraph(document)
    add_paragraph(document, "Вариант: 3")
    add_paragraph(
        document,
        "Методы: сопряжённые градиенты Флетчера–Ривза, "
        "модифицированный метод Ньютона, BFGS.",
    )
    add_paragraph(
        document,
        "Цель работы: реализовать и сравнить методы безусловной "
        "минимизации на тестовых функциях и полиноме 10D.",
    )

    document.add_page_break()


def add_problem_statement(document: Document, epsilon: float) -> None:
    add_heading(document, "1. Постановка задачи", level=1)

    add_paragraph(
        document,
        "Необходимо решить задачу безусловной минимизации функции f(x) "
        "тремя методами: методом сопряжённых градиентов, модифицированным "
        "методом Ньютона и методом BFGS.",
    )

    add_paragraph(
        document,
        f"Критерий остановки: ||∇f(xᵏ)|| ≤ ε, где ε = {epsilon:g}. "
        "Длина шага выбирается с помощью бэктрекинга Армихо.",
    )

    add_paragraph(
        document,
        "Для двумерных тестовых функций используются аналитические "
        "градиенты и Гессианы. Для полинома 10D в программе также "
        "реализованы аналитический градиент и аналитический Гессиан.",
    )


def add_variant_description(document: Document) -> None:
    add_heading(document, "2. Функции варианта 3", level=1)

    dataframe = pd.DataFrame(
        [
            {
                "№": 1,
                "Функция": "Бут",
                "Начальная точка": "(-5, -5)",
            },
            {
                "№": 2,
                "Функция": "Сфера",
                "Начальная точка": "(-3, -3)",
            },
            {
                "№": 3,
                "Функция": "Розенброк",
                "Начальная точка": "(-2, 2)",
            },
            {
                "№": 4,
                "Функция": "Химмельблау",
                "Начальная точка": "(-1, 0)",
            },
            {
                "№": 5,
                "Функция": "Полином 10D",
                "Начальная точка": "(-1.5, -2.4, 2.5, 0, ..., 0)",
            },
        ]
    )

    add_dataframe_table(document, dataframe)


def add_method_summary(document: Document) -> None:
    add_heading(document, "3. Краткое описание методов", level=1)

    add_heading(document, "3.1. Метод сопряжённых градиентов Флетчера–Ривза", level=2)
    add_paragraph(
        document,
        "В методе используется направление p₀ = -g₀, а последующие "
        "направления строятся по формуле pₖ = -gₖ + βₖpₖ₋₁. "
        "Для варианта Флетчера–Ривза βₖ = ||gₖ||² / ||gₖ₋₁||².",
    )

    add_heading(document, "3.2. Модифицированный метод Ньютона", level=2)
    add_paragraph(
        document,
        "На каждой итерации решается система Hₖpₖ = -gₖ. "
        "Если Гессиан не является положительно определённым, "
        "используется регуляризация Hₖ + μI.",
    )

    add_heading(document, "3.3. Метод BFGS", level=2)
    add_paragraph(
        document,
        "BFGS является квазиньютоновским методом. В реализации хранится "
        "приближение обратного Гессиана, которое обновляется по формуле BFGS "
        "при выполнении условия кривизны yₖᵀsₖ > 0.",
    )


def add_2d_function_section(
    document: Document,
    function: ObjectiveFunction,
    function_results: list[OptimizationResult],
    epsilon: float,
) -> None:
    function_name = get_function_ru_name(function.name)

    add_heading(document, f"4. Результаты для функции {function_name}", level=1)

    add_paragraph(
        document,
        f"Начальная точка: x⁰ = {format_array(function.default_x0)}.",
    )

    if function.known_minima:
        minima_text = "; ".join(
            format_array(point)
            for point in function.known_minima
        )
        add_paragraph(document, f"Известные минимумы: {minima_text}.")

    add_heading(document, "4.1. Сводная таблица результатов", level=2)
    function_summary = prepare_function_summary_dataframe(function_results)
    add_dataframe_table(document, function_summary)

    add_heading(document, "4.2. График линий уровня и траекторий", level=2)
    contour_path = plot_contour_with_tracks(function, function_results)
    add_image_if_exists(document, contour_path)

    add_heading(document, "4.3. График сходимости", level=2)
    convergence_path = plot_gradient_norm_convergence(
        function=function,
        results=function_results,
        epsilon=epsilon,
    )
    add_image_if_exists(document, convergence_path)

    add_heading(document, "4.4. График убывания значения функции", level=2)
    decrease_path = plot_function_decrease(function, function_results)
    add_image_if_exists(document, decrease_path)

    add_heading(document, "4.5. Таблицы итераций", level=2)

    for result in function_results:
        add_heading(document, get_method_ru_name(result.method_name), level=3)
        history_dataframe = prepare_history_dataframe(result)
        add_dataframe_table(document, history_dataframe)

    add_heading(document, "4.6. Вывод", level=2)
    best_result = min(function_results, key=lambda item: item.f_star)
    add_paragraph(
        document,
        f"Наименьшее значение функции получено методом "
        f"{get_method_ru_name(best_result.method_name)}: "
        f"f(x*) ≈ {best_result.f_star:.6g}, "
        f"||∇f(x*)|| ≈ {best_result.grad_norm:.6g}.",
    )


def add_polynomial_section(
    document: Document,
    function: ObjectiveFunction,
    function_results: list[OptimizationResult],
    epsilon: float,
) -> None:
    function_name = get_function_ru_name(function.name)

    add_heading(document, f"5. Результаты для {function_name}", level=1)

    add_paragraph(
        document,
        f"Начальная точка: x⁰ = {format_array(function.default_x0)}.",
    )

    add_heading(document, "5.1. Сводная таблица результатов", level=2)
    function_summary = prepare_function_summary_dataframe(function_results)
    add_dataframe_table(document, function_summary)

    add_heading(document, "5.2. Найденные точки x*", level=2)

    rows = []
    for result in function_results:
        rows.append(
            {
                "Метод": get_method_ru_name(result.method_name),
                "x*": format_array(result.x_star),
                "f(x*)": result.f_star,
                "||∇f(x*)||": result.grad_norm,
            }
        )

    add_dataframe_table(document, pd.DataFrame(rows))

    add_heading(document, "5.3. График сходимости", level=2)
    convergence_path = plot_gradient_norm_convergence(
        function=function,
        results=function_results,
        epsilon=epsilon,
    )
    add_image_if_exists(document, convergence_path)

    add_heading(document, "5.4. График убывания значения функции", level=2)
    decrease_path = plot_function_decrease(function, function_results)
    add_image_if_exists(document, decrease_path)

    add_heading(document, "5.5. 2D-сечение полинома", level=2)
    section_path = plot_polynomial_2d_section(
        function=function,
        results=function_results,
        variable_i=0,
        variable_j=1,
    )
    add_image_if_exists(document, section_path)

    add_heading(document, "5.6. Таблицы итераций", level=2)

    for result in function_results:
        add_heading(document, get_method_ru_name(result.method_name), level=3)
        history_dataframe = prepare_history_dataframe(result)
        add_dataframe_table(document, history_dataframe)

    add_heading(document, "5.7. Вывод", level=2)
    best_result = min(function_results, key=lambda item: item.f_star)
    add_paragraph(
        document,
        f"Лучшее значение для полинома получено методом "
        f"{get_method_ru_name(best_result.method_name)}: "
        f"f(x*) ≈ {best_result.f_star:.6g}, "
        f"||∇f(x*)|| ≈ {best_result.grad_norm:.6g}.",
    )


def add_global_summary(
    document: Document,
    results: list[OptimizationResult],
) -> None:
    add_heading(document, "6. Сводная таблица сравнения", level=1)

    summary_dataframe = prepare_summary_dataframe(results)
    add_dataframe_table(document, summary_dataframe)


def add_final_conclusion(
    document: Document,
    results: list[OptimizationResult],
) -> None:
    add_heading(document, "7. Общий вывод", level=1)

    add_paragraph(
        document,
        "Все три обязательных метода были реализованы и применены к функциям "
        "варианта 3. Для всех двумерных функций построены линии уровня с "
        "траекториями методов, а также графики сходимости и убывания значения "
        "целевой функции.",
    )

    add_paragraph(
        document,
        "Модифицированный метод Ньютона обычно требует меньше итераций, "
        "поскольку использует информацию второго порядка. Однако одна итерация "
        "метода Ньютона дороже из-за вычисления и обработки Гессиана.",
    )

    add_paragraph(
        document,
        "Метод BFGS показывает устойчивую сходимость и часто требует число "
        "итераций, близкое к методу Ньютона, но не требует явного вычисления "
        "Гессиана.",
    )

    add_paragraph(
        document,
        "Метод сопряжённых градиентов имеет простую структуру и низкую стоимость "
        "итерации, но на овражных функциях, например на функции Розенброка, "
        "может требовать существенно больше итераций.",
    )

    converged_count = sum(result.converged for result in results)
    add_paragraph(
        document,
        f"Количество успешных запусков: {converged_count} из {len(results)}.",
    )


def configure_document_styles(document: Document) -> None:
    styles = document.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.bold = True


def generate_docx_report(
    functions: list[ObjectiveFunction],
    results: list[OptimizationResult],
    epsilon: float,
    output_path: Path | None = None,
) -> Path:
    """
    Создание DOCX-отчёта по результатам лабораторной работы.

    Отчёт содержит:
    - постановку задачи;
    - описание варианта;
    - описание методов;
    - таблицы результатов;
    - графики;
    - таблицы итераций;
    - общий вывод.
    """

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)

    if output_path is None:
        output_path = REPORTS_DIR / "report_variant_3.docx"

    document = Document()
    configure_document_styles(document)

    add_title_page(document)
    add_problem_statement(document, epsilon)
    add_variant_description(document)
    add_method_summary(document)

    for function in functions:
        function_results = get_results_for_function(function, results)

        if not function_results:
            continue

        if function.dimension == 2:
            add_2d_function_section(
                document=document,
                function=function,
                function_results=function_results,
                epsilon=epsilon,
            )
        elif function.dimension == 10:
            add_polynomial_section(
                document=document,
                function=function,
                function_results=function_results,
                epsilon=epsilon,
            )

    add_global_summary(document, results)
    add_final_conclusion(document, results)

    document.save(output_path)

    return output_path